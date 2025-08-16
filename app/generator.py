# app/generator.py
from docx import Document
import os
import subprocess
import shutil
import platform


# Detect system and set LibreOffice path dynamically
if platform.system() == "Windows":
    LIBREOFFICE_PATH = r"C:\Program Files\LibreOffice\program\soffice.exe"
else:
    # On Linux (Render container), soffice will be installed from apt-get
    LIBREOFFICE_PATH = shutil.which("soffice") or "/usr/bin/soffice"


def replace_text_in_docx(input_docx_path, output_docx_path, replacements):
    """Replace placeholders in DOCX template and save new DOCX."""
    doc = Document(input_docx_path)
    for paragraph in doc.paragraphs:
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(old, new)
    doc.save(output_docx_path)


def convert_docx_to_pdf(input_docx_path, output_pdf_path):
    """
    Convert DOCX to PDF using LibreOffice in headless mode.
    Works both on Windows and Linux (Render).
    """
    if not LIBREOFFICE_PATH or not os.path.exists(LIBREOFFICE_PATH):
        raise FileNotFoundError(f"LibreOffice not found at: {LIBREOFFICE_PATH}")

    output_dir = os.path.dirname(output_pdf_path)
    os.makedirs(output_dir, exist_ok=True)

    # Run LibreOffice in headless mode
    subprocess.run([
        LIBREOFFICE_PATH, "--headless", "--convert-to", "pdf", "--outdir", output_dir, input_docx_path
    ], check=True)

    # LibreOffice generates file with original name but .pdf extension
    generated_pdf = os.path.join(
        output_dir,
        os.path.splitext(os.path.basename(input_docx_path))[0] + ".pdf"
    )

    if not os.path.exists(generated_pdf):
        raise FileNotFoundError(f"PDF was not generated for {input_docx_path}")

    # Rename to desired output path
    os.rename(generated_pdf, output_pdf_path)


def generate_certificate(row, template_file, output_dir='.'):
    """Generate certificate from template with replaced placeholders."""
    row_name = row['Name'].strip()
    row_domain = row['Domain']

    replacements = {
        '<NAME>': row_name,
        '<DOMAIN>': row_domain,
    }

    os.makedirs(output_dir, exist_ok=True)

    output_docx = os.path.join(output_dir, f'{row_name}_Certificate.docx')
    output_pdf = os.path.join(output_dir, f'DNYX-Completion-{row_name}.pdf')

    replace_text_in_docx(template_file, output_docx, replacements)
    print(f"[INFO] DOCX created: {output_docx}")

    convert_docx_to_pdf(output_docx, output_pdf)
    print(f"[INFO] PDF created: {output_pdf}")

    os.remove(output_docx)  # Remove intermediate DOCX file
    return output_pdf


def generate_offer_letter(row, template_file, output_dir='.'):
    """Generate offer letter from template with replaced placeholders."""
    row_name = row['Name'].strip()
    row_domain = row['Domain']

    replacements = {
        '<NAME>': row_name,
        '<DOMAIN>': row_domain,
    }

    os.makedirs(output_dir, exist_ok=True)

    output_docx = os.path.join(output_dir, f'{row_name}_OfferLetter.docx')
    output_pdf = os.path.join(output_dir, f'DNYX-OfferLetter-{row_name}.pdf')

    replace_text_in_docx(template_file, output_docx, replacements)
    print(f"[INFO] DOCX created: {output_docx}")

    convert_docx_to_pdf(output_docx, output_pdf)
    print(f"[INFO] PDF created: {output_pdf}")

    os.remove(output_docx)  # Remove intermediate DOCX file
    return output_pdf
